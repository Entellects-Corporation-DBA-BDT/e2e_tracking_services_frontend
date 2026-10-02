from pathlib import Path
import fitz,json,math
from openpyxl import load_workbook
R=Path(__file__).parent;O=R/"Navarasa_Kitchen_Premium_Menu.pdf"; W,H=612,792
M=(.38,.055,.06);G=(.78,.55,.2);C=(.985,.965,.91);I=(.16,.12,.095);U=(.37,.32,.27)
fs={"r":r"C:\Windows\Fonts\georgia.ttf","b":r"C:\Windows\Fonts\georgiab.ttf","s":r"C:\Windows\Fonts\segoeui.ttf","sb":r"C:\Windows\Fonts\segoeuib.ttf"}
w=load_workbook(R/"menu.xlsx",data_only=True).active
A=[(str(a),str(b),str(c),str(d or "")) for a,b,c,d in w.iter_rows(min_row=4,values_only=True) if a]
D=fitz.open()
def setup(p):
 for n,f in fs.items():p.insert_font(fontname=n,fontfile=f)
def image(p,r,f,k=True):p.insert_image(r,filename=str(R/f),keep_proportion=k)
def wrap(t,f,z,m):
 q="";a=[]
 for v in t.split():
  n=(q+" "+v).strip()
  if fitz.get_text_length(n,fontname=("Times-Bold" if f=="b" else "Helvetica"),fontsize=z)<=m:q=n
  else:
   if q:a.append(q)
   q=""
   for ch in v:
    if fitz.get_text_length(q+ch,fontname=("Times-Bold" if f=="b" else "Helvetica"),fontsize=z)<=m:q+=ch
    else:a.append(q);q=ch
 if q:a.append(q)
 return a
def base(p,n):
 p.draw_rect(p.rect,fill=C,color=None);p.draw_rect((0,0,W,72),fill=M,color=None);p.draw_rect((0,72,W,76),fill=G,color=None)
 p.draw_circle((306,405),104,color=G,width=.7,stroke_opacity=.07);p.insert_textbox((220,350,392,455),"N",fontname="b",fontsize=86,color=G,align=1,fill_opacity=.05)
 p.insert_text((36,34),"NAVARASA",fontname="b",fontsize=20,color=(1,1,1));p.insert_text((36,53),"K I T C H E N",fontname="s",fontsize=8.5,color=(.94,.87,.68))
 p.insert_textbox((330,24,576,56),"TASTE THE INDIAN TRADITION",fontname="sb",fontsize=8.5,color=(.94,.87,.68),align=2)
 p.draw_line((36,750),(576,750),color=G);p.insert_text((36,769),"Allergy Notice: Please inform our staff of any food allergies or dietary requirements before ordering.",fontname="s",fontsize=5.8,color=U)
 p.insert_text((36,778),"Our kitchen handles common allergens, and cross-contact may occur.",fontname="s",fontsize=5.8,color=U);p.insert_textbox((540,758,576,778),str(n),fontname="sb",fontsize=7,color=M,align=2)
def new(n):
 p=D.new_page(width=W,height=H);setup(p);base(p,n);return p
def met(a):
 c,d,pr,de=a;pr=pr.replace("🌶️","").strip();nl=wrap(d,"b",8.7,176);dl=wrap(de,"s",6.5,244);return nl,dl,max(12,len(nl)*10)+len(dl)*7.6+8,pr,"🌶" in a[2]
def hd(p,x,y,t,cont):
 p.draw_rect((x,y,x+262,y+24),fill=M,color=None);p.draw_rect((x,y+21,x+262,y+24),fill=G,color=None);p.insert_textbox((x+10,y+4,x+254,y+20),t+(" · CONTINUED" if cont else ""),fontname="sb",fontsize=8.3,color=(1,1,1));return y+31
def it(p,x,y,a):
 nl,dl,h,pr,hot=met(a)
 for j,l in enumerate(nl):p.insert_text((x+8,y+9+j*10),l,fontname="b",fontsize=8.7,color=I)
 pw=fitz.get_text_length(pr,fontname="Helvetica-Bold",fontsize=8.8);p.insert_text((x+254-pw,y+9),pr,fontname="sb",fontsize=8.8,color=M)
 if hot:p.draw_circle((x+247-pw,y+5),2.3,fill=(.7,.08,.04),color=(.7,.08,.04))
 yy=y+max(12,len(nl)*10)+1
 for l in dl:p.insert_text((x+8,yy+6),l,fontname="s",fontsize=6.5,color=U);yy+=7.6
 p.draw_line((x+8,y+h-3),(x+254,y+h-3),color=G,width=.25,stroke_opacity=.45);return y+h
p=D.new_page(width=W,height=H);setup(p);p.draw_rect(p.rect,fill=(.12,.09,.075),color=None);image(p,(0,0,W,350),"indian_feast.png",False)
for i in range(12):p.draw_rect((0,245+i*10,W,257+i*10),fill=(.12,.09,.075),color=None,fill_opacity=i/15)
p.draw_rect((24,24,588,768),color=G,width=1.2);p.draw_circle((306,265),88,fill=(1,1,1),color=G,width=2);image(p,(223,179,389,351),"logo.jpeg")
p.insert_textbox((50,376,562,435),"NAVARASA",fontname="b",fontsize=42,color=(1,1,1),align=1);p.insert_textbox((50,430,562,468),"K I T C H E N",fontname="sb",fontsize=16,color=G,align=1)
p.insert_textbox((60,505,552,540),"Taste the Indian Tradition",fontname="r",fontsize=17,color=(.94,.87,.68),align=1);p.insert_textbox((90,565,522,615),"A celebration of regional Indian flavors, crafted with warmth and served with heart.",fontname="s",fontsize=10,color=(.93,.89,.8),align=1)
image(p,(266,636,346,710),"halal_clean.png");p.insert_textbox((50,721,562,744),"PLAINFIELD, ILLINOIS  •  MENU",fontname="sb",fontsize=8.5,color=G,align=1)
xs=[36,314];pn=2;p=new(pn);co=0;y=96;cat="";first=True
for a in A:
 if a[0]!=cat:cat=a[0];first=True
 need=met(a)[2]+(31 if first else 0)
 if y+need>742:
  co+=1
  if co==2:pn+=1;p=new(pn);co=0
  y=hd(p,xs[co],96,cat,not first);first=False
 elif first:y=hd(p,xs[co],y,cat,False);first=False
 y=it(p,xs[co],y,a)
D.set_metadata({"title":"Navarasa Kitchen — Premium Menu","author":"Navarasa Kitchen"});D.save(O,garbage=4,deflate=True,clean=True);D.close()
q=fitz.open(O);t="\n".join(p.get_text() for p in q);rep={"items":len(A),"categories":len(dict.fromkeys(x[0] for x in A)),"pages":q.page_count,"missing_dishes":[x[1] for x in A if x[1] not in t],"missing_prices":[x[1] for x in A if x[2].replace("🌶️","").strip() not in t],"tagline":"Taste the Indian Tradition" in t,"allergy":"Allergy Notice:" in t};(R/"validation_report.json").write_text(json.dumps(rep,indent=2),encoding="utf8");print(json.dumps(rep,indent=2))