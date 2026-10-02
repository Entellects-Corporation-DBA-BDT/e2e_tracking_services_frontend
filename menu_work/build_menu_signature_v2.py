from pathlib import Path
import fitz,json,math
from openpyxl import load_workbook
R=Path(__file__).parent;O=R/"Navarasa_Kitchen_Signature_Menu_v2.pdf"; W,H=612,792
M=(.035,.29,.28);G=(.92,.46,.12);C=(1,.97,.91);I=(.20,.08,.16);U=(.34,.29,.27)
fs={"r":r"C:\Windows\Fonts\georgia.ttf","b":r"C:\Windows\Fonts\georgiab.ttf","s":r"C:\Windows\Fonts\segoeui.ttf","sb":r"C:\Windows\Fonts\segoeuib.ttf"}
w=load_workbook(R/"menu.xlsx",data_only=True).active
A=[(str(a),str(b),str(c),str(d or "")) for a,b,c,d in w.iter_rows(min_row=4,values_only=True) if a]
cats=list(dict.fromkeys(x[0] for x in A))
D=fitz.open()
def setup(p):
 for n,f in fs.items():p.insert_font(fontname=n,fontfile=f)
def image(p,r,f,k=True):p.insert_image(r,filename=str(R/f),keep_proportion=k)
def wrap(t,f,z,m):
 q="";a=[]
 for v in t.split():
  n=(q+" "+v).strip()
  if fitz.get_text_length(n,fontname=("Times-Bold" if f=="b" else "Helvetica"),fontsize=z)<=m:q=n
  else:
   if q:a.append(q)
   q=""
   for ch in v:
    if fitz.get_text_length(q+ch,fontname=("Times-Bold" if f=="b" else "Helvetica"),fontsize=z)<=m:q+=ch
    else:a.append(q);q=ch
 if q:a.append(q)
 return a
def base(p,n):
 p.draw_rect(p.rect,fill=C,color=None)
 for xx in range(12,612,36):
  for yy in range(96,744,36):
   p.draw_circle((xx,yy),7,color=G,width=.25,stroke_opacity=.08);p.draw_line((xx-7,yy),(xx,yy-7),color=M,width=.22,stroke_opacity=.06);p.draw_line((xx,yy-7),(xx+7,yy),color=M,width=.22,stroke_opacity=.06)
 image(p,(198,292,414,508),"logo_watermark.png")
 p.draw_rect((0,0,W,82),fill=M,color=None);p.draw_rect((0,0,W,5),fill=G,color=None);p.draw_rect((0,70,W,76),fill=G,color=None)
 for xx in range(12,612,24):p.draw_circle((xx,76),6,fill=G,color=G)
 p.draw_circle((48,40),34,fill=(1,.97,.91),color=G,width=1.5);image(p,(17,9,79,71),"logo_transparent.png")
 p.insert_text((94,32),"NAVARASA KITCHEN",fontname="b",fontsize=19,color=(1,1,1));p.insert_text((94,53),"A CELEBRATION OF INDIAN TRADITION",fontname="s",fontsize=7.4,color=(1,.88,.58))
 p.draw_line((368,25),(368,58),color=G,width=.6);p.insert_textbox((386,24,574,44),"PREMIUM DINING MENU",fontname="sb",fontsize=8.2,color=(1,1,1),align=2);p.insert_textbox((386,45,574,59),"PLAINFIELD, ILLINOIS",fontname="s",fontsize=6.8,color=(1,.87,.56),align=2)
 p.draw_line((36,750),(576,750),color=G,width=1);p.insert_text((36,769),"Allergy Notice: Please inform our staff of any food allergies or dietary requirements before ordering.",fontname="s",fontsize=5.8,color=U);p.insert_text((36,778),"Our kitchen handles common allergens, and cross-contact may occur.",fontname="s",fontsize=5.8,color=U)
 p.draw_circle((556,768),14,fill=M,color=G,width=1);p.insert_textbox((542,762,570,775),str(n),fontname="sb",fontsize=7,color=(1,1,1),align=1)
def new(n):
 p=D.new_page(width=W,height=H);setup(p);base(p,n);return p
def met(a):
 c,d,pr,de=a;pr=pr.replace("ðŸŒ¶ï¸","").strip();nl=wrap(d,"b",8.7,176);dl=wrap(de,"s",6.5,244);return nl,dl,max(12,len(nl)*10)+len(dl)*7.6+8,pr,"ðŸŒ¶" in a[2]
def hd(p,x,y,t,cont):
 p.draw_rect((x+20,y+7,x+262,y+43),fill=(.88,.73,.48),color=None,fill_opacity=.22);p.draw_rect((x+16,y+3,x+258,y+39),fill=(1,.955,.85),color=G,width=1.1)
 p.draw_line((x+50,y+38),(x+250,y+38),color=M,width=.7);p.draw_circle((x+25,y+21),25,fill=C,color=G,width=2.2)
 idx=cats.index(t) if t in cats else 19;image(p,(x+2,y-2,x+48,y+44),f"circle_{idx:02}.png",False)
 p.insert_textbox((x+56,y+10,x+248,y+31),t+(" Â· CONTINUED" if cont else ""),fontname="b",fontsize=8.6,color=M);p.draw_circle((x+250,y+21),3,fill=G,color=None)
 return y+54
def it(p,x,y,a):
 nl,dl,h,pr,hot=met(a)
 for j,l in enumerate(nl):p.insert_text((x+8,y+9+j*10),l,fontname="b",fontsize=8.7,color=I)
 pw=fitz.get_text_length(pr,fontname="Helvetica-Bold",fontsize=8.8);p.insert_text((x+254-pw,y+9),pr,fontname="sb",fontsize=8.8,color=M)
 if hot:p.draw_circle((x+247-pw,y+5),2.3,fill=(.7,.08,.04),color=(.7,.08,.04))
 yy=y+max(12,len(nl)*10)+1
 for l in dl:p.insert_text((x+8,yy+6),l,fontname="s",fontsize=6.5,color=U);yy+=7.6
 p.draw_line((x+8,y+h-3),(x+254,y+h-3),color=G,width=.25,stroke_opacity=.45);return y+h
p=D.new_page(width=W,height=H);setup(p);p.draw_rect(p.rect,fill=C,color=None)
for xx in range(18,612,42):
 for yy in range(18,792,42):p.draw_circle((xx,yy),8,color=G,width=.3,stroke_opacity=.10)
p.draw_rect((0,0,W,286),fill=M,color=None);image(p,(0,0,W,286),"indian_feast.png",False);p.draw_rect((0,205,W,286),fill=M,color=None,fill_opacity=.56)
p.draw_rect((18,18,594,774),color=G,width=1.6);p.draw_rect((25,25,587,767),color=M,width=.45)
for rr in (102,94,86):p.draw_circle((306,282),rr,color=G,width=1.2 if rr==102 else .5,fill=C if rr==86 else None)
for a in range(0,360,30):
 rad=math.radians(a);p.draw_circle((306+111*math.cos(rad),282+111*math.sin(rad)),4,fill=G,color=None)
image(p,(218,194,394,370),"logo_transparent.png")
p.insert_textbox((60,402,552,452),"NAVARASA KITCHEN",fontname="b",fontsize=31,color=M,align=1);p.insert_textbox((100,452,512,480),"Taste the Indian Tradition",fontname="r",fontsize=16,color=G,align=1)
p.draw_line((142,494),(470,494),color=G,width=1.2);p.insert_textbox((78,514,534,558),"An elevated journey through the regional flavors, aromas and hospitality of India.",fontname="s",fontsize=10.5,color=I,align=1)
labels=["AUTHENTIC FLAVORS","REGIONAL FAVORITES","HALAL"]
for i,l in enumerate(labels):
 x=74+i*158;p.draw_rect((x,588,x+142,620),fill=M,color=G,width=.8);p.insert_textbox((x+4,598,x+138,614),l,fontname="sb",fontsize=7.3,color=(1,.91,.62),align=1)
image(p,(265,642,347,718),"halal_clean.png");p.insert_textbox((60,730,552,751),"PLAINFIELD, ILLINOIS  â€¢  PREMIUM DINING MENU",fontname="sb",fontsize=8.5,color=M,align=1)# Visual guide / kitchen story page
p=new(2);image(p,(36,96,576,330),"chef_kitchen.png",False)
p.draw_rect((36,300,576,330),fill=M,color=None,fill_opacity=.88);p.insert_textbox((50,306,562,326),"FROM OUR KITCHEN TO YOUR TABLE",fontname="sb",fontsize=12,color=(1,1,1),align=1)
p.insert_textbox((60,348,552,390),"Explore our menu by category",fontname="b",fontsize=19,color=M,align=1)
for i,c in enumerate(cats):
 col=i%4;row=i//4;x=44+col*134;y0=410+row*60
 image(p,(x,y0,x+42,y0+42),f"circle_{i:02}.png",False)
 p.insert_textbox((x+48,y0+5,x+126,y0+40),c,fontname="sb",fontsize=6.8,color=I)
p.insert_textbox((60,712,552,738),"Every dish and price on the following pages is presented exactly from our current menu.",fontname="s",fontsize=8,color=U,align=1)
xs=[36,314];pn=3;p=new(pn);co=0;y=96;cat="";first=True
for a in A:
 if a[0]!=cat:cat=a[0];first=True
 need=met(a)[2]+(54 if first else 0)
 if y+need>742:
  co+=1
  if co==2:pn+=1;p=new(pn);co=0
  y=hd(p,xs[co],96,cat,not first);first=False
 elif first:y=hd(p,xs[co],y,cat,False);first=False
 y=it(p,xs[co],y,a)
D.set_metadata({"title":"Navarasa Kitchen â€” Premium Menu","author":"Navarasa Kitchen"});D.save(O,garbage=4,deflate=True,clean=True);D.close()
q=fitz.open(O);t="\n".join(p.get_text() for p in q);rep={"items":len(A),"categories":len(dict.fromkeys(x[0] for x in A)),"pages":q.page_count,"missing_dishes":[x[1] for x in A if x[1] not in t],"missing_prices":[x[1] for x in A if x[2].replace("ðŸŒ¶ï¸","").strip() not in t],"tagline":"Taste the Indian Tradition" in t,"allergy":"Allergy Notice:" in t};(R/"validation_report.json").write_text(json.dumps(rep,indent=2),encoding="utf8");print(json.dumps(rep,indent=2))
